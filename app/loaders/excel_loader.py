from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.domain import ExtractedChunk


def load_excel(path: Path, rows_per_chunk: int) -> list[ExtractedChunk]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    chunks: list[ExtractedChunk] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            start_row = 1
            max_col = 1
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = ["" if value is None else str(value) for value in row]
                max_col = max(max_col, len(values))
                rows.append("\t".join(values))
                if len(rows) >= rows_per_chunk:
                    chunks.append(ExtractedChunk(text="\n".join(rows), chunk_type="excel_range", sheet_name=sheet.title, cell_range=f"A{start_row}:{get_column_letter(max_col)}{row_number}"))
                    rows, start_row = [], row_number + 1
            if rows:
                end_row = start_row + len(rows) - 1
                chunks.append(ExtractedChunk(text="\n".join(rows), chunk_type="excel_range", sheet_name=sheet.title, cell_range=f"A{start_row}:{get_column_letter(max_col)}{end_row}"))
    finally:
        workbook.close()
    return chunks
